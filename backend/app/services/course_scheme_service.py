"""Course Scheme services: file text extraction, Supabase Storage, and Course materialization.

Text extraction tries the real text layer first and only falls back to OCR when
a PDF turns out to be a scan. Storage keeps the original upload so the admin can
always go back to the source document. See docs/PROJECT_ARCHITECTURE.md §7.
"""

import io
import re
import uuid
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any

import httpx
import pdfplumber
from sqlalchemy import delete, func, inspect, select, text
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Course, CourseScheme

SCHEME_BUCKET = "course-schemes"

# A real text PDF gives us far more than this per page; anything less means the
# page is an image and needs OCR.
MIN_CHARS_PER_PAGE = 120

# Scheme documents list a lab as its own row, e.g. "PROGRAMMING FUNDAMENTALS (LAB)".
LAB_SUFFIX_RE = re.compile(r"\s*\(\s*LAB\s*\)\s*$", re.IGNORECASE)

OCR_DPI = 300
STORAGE_TIMEOUT_SECONDS = 60

INSTALL_HINT = (
    "OCR needs the Tesseract binary and poppler on PATH. On Windows: "
    "`winget install --id UB-Mannheim.TesseractOCR` and `winget install --id oschwartz10612.Poppler`, "
    "then reopen the terminal."
)


class SchemeFileError(Exception):
    """Raised when an uploaded file cannot be read; the message is shown to the admin."""


class OcrUnavailableError(SchemeFileError):
    """Raised when a scanned file needs OCR but the OCR tooling is not installed."""


class StorageError(Exception):
    """Raised when Supabase Storage rejects an upload or delete."""


@dataclass
class ExtractionResult:
    text: str
    method: str  # "pdf-text", "ocr" or "word"
    page_count: int
    warnings: list[str] = field(default_factory=list)


def extract_text(filename: str, data: bytes) -> ExtractionResult:
    # Entry point for an upload: picks the reader based on the file extension.
    suffix = PurePath(filename).suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix in {".docx", ".doc"}:
        return _extract_word(suffix, data)

    raise SchemeFileError(f"Unsupported file type '{suffix or filename}'. Upload a PDF or Word (.docx) file.")


def _extract_pdf(data: bytes) -> ExtractionResult:
    # Reads the PDF's own text layer, and only rasterizes for OCR if that comes back too thin.
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        page_count = len(pdf.pages)
        text = "\n\n".join(page.extract_text() or "" for page in pdf.pages).strip()

    if not _looks_scanned(text, page_count):
        return ExtractionResult(text=text, method="pdf-text", page_count=page_count)

    ocr_text = _ocr_pdf(data)
    warning = (
        "This PDF has no usable text layer, so it was read with OCR. "
        "Check the text carefully — OCR misreads characters."
    )
    return ExtractionResult(text=ocr_text, method="ocr", page_count=page_count, warnings=[warning])


def _looks_scanned(text: str, page_count: int) -> bool:
    # A scan yields almost no characters per page, which is how we tell it apart from a real text PDF.
    if page_count == 0:
        return False
    return len(text.strip()) / page_count < MIN_CHARS_PER_PAGE


def _ocr_pdf(data: bytes) -> str:
    # Renders each page to an image and runs Tesseract over it. Both poppler and the
    # tesseract binary must be installed, so missing tooling is reported precisely.
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except ImportError as exc:  # pragma: no cover - packages are in requirements
        raise OcrUnavailableError(f"OCR packages are not installed. {INSTALL_HINT}") from exc

    try:
        pages = convert_from_bytes(data, dpi=OCR_DPI)
    except Exception as exc:  # pdf2image raises its own error type when poppler is absent
        raise OcrUnavailableError(f"Could not rasterize the PDF for OCR. {INSTALL_HINT}") from exc

    try:
        return "\n\n".join(pytesseract.image_to_string(page) for page in pages).strip()
    except pytesseract.TesseractNotFoundError as exc:
        raise OcrUnavailableError(f"Tesseract is not installed. {INSTALL_HINT}") from exc


def _extract_word(suffix: str, data: bytes) -> ExtractionResult:
    # Pulls paragraphs and table cells out of a .docx; legacy .doc is not a zip archive and can't be read.
    if suffix == ".doc":
        raise SchemeFileError("Legacy .doc files are not supported. Save the file as .docx or PDF and retry.")

    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return ExtractionResult(text="\n".join(lines), method="word", page_count=0)


def _normalize_name(name: str) -> str:
    # Case- and spacing-insensitive form of a course name, used to match a lab to its theory course.
    return " ".join(name.split()).upper()


def _lab_base_name(name: str) -> str | None:
    # For a "NAME (LAB)" row, the normalized theory name it belongs to; None for an ordinary course.
    match = LAB_SUFFIX_RE.search(name)
    return _normalize_name(name[: match.start()]) if match else None


def pair_lab_rows(courses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    # Folds each "(LAB)" row into the theory course with the same name, so a lab becomes
    # has_lab + lab_credit_hours on that course instead of a course of its own. Order is kept.
    # A lab with no matching theory course stays as its own row rather than being dropped.
    rows = [dict(course) for course in courses]

    theory_by_name: dict[str, dict[str, Any]] = {}
    for row in rows:
        if _lab_base_name(row["name"]) is None:
            theory_by_name.setdefault(_normalize_name(row["name"]), row)

    paired = []
    for row in rows:
        base_name = _lab_base_name(row["name"])
        if base_name is None:
            paired.append(row)
        elif base_name in theory_by_name:
            theory = theory_by_name[base_name]
            theory["has_lab"] = True
            theory["lab_credit_hours"] = row.get("credit_hours")
        else:
            row["has_lab"] = True
            row["lab_credit_hours"] = row.get("lab_credit_hours") or row.get("credit_hours")
            paired.append(row)
    return paired


def paired_semesters(content: dict[str, Any]) -> list[dict[str, Any]]:
    # The scheme's semesters with labs folded into their theory courses (pairing never crosses
    # semesters). This is exactly what becomes Course rows, and what the UI shows for a saved scheme.
    return [
        {"semester": semester.get("semester"), "courses": pair_lab_rows(semester.get("courses", []))}
        for semester in content.get("semesters", [])
    ]


def materialize_courses(session: Session, scheme: CourseScheme) -> int:
    # Rebuilds the queryable Course rows from the scheme's content JSON, which stays the source of record.
    session.execute(delete(Course).where(Course.scheme_id == scheme.id))

    rows = [
        Course(
            scheme_id=scheme.id,
            code=course["code"],
            name=course["name"],
            credit_hours=course.get("credit_hours"),
            has_lab=bool(course.get("has_lab", False)),
            lab_credit_hours=course.get("lab_credit_hours"),
            min_marks=course.get("min_marks"),
            max_marks=course.get("max_marks"),
        )
        for semester in paired_semesters(scheme.content)
        for course in semester["courses"]
    ]
    session.add_all(rows)
    return len(rows)


def course_counts(session: Session, scheme_ids: list[int]) -> dict[int, tuple[int, int]]:
    # Total and lab-bearing Course counts per scheme, in one grouped query instead of one per scheme.
    if not scheme_ids:
        return {}

    rows = session.execute(
        select(
            Course.scheme_id,
            func.count(Course.id),
            func.count(Course.id).filter(Course.has_lab.is_(True)),
        )
        .where(Course.scheme_id.in_(scheme_ids))
        .group_by(Course.scheme_id)
    ).all()
    return {scheme_id: (total, labs) for scheme_id, total, labs in rows}


def find_blocking_timetables(session: Session, scheme_id: int) -> list[str]:
    # Deleting a scheme must never orphan a published timetable. Division and Timetable models
    # arrive in a later phase, so the query runs against raw SQL and only once those tables exist;
    # until then nothing can reference a scheme and the answer is genuinely "nothing blocks it".
    tables = _existing_table_names(session)
    if not {"timetables", "divisions"} <= tables:
        return []

    rows = session.execute(
        text(
            "SELECT t.id FROM timetables t "
            "JOIN divisions d ON d.id = t.division_id "
            "WHERE d.scheme_year_id = :scheme_id AND t.status = 'published'"
        ),
        {"scheme_id": scheme_id},
    ).scalars().all()
    return [f"timetable #{row}" for row in rows]


def _existing_table_names(session: Session) -> set[str]:
    # Looks at what tables actually exist, so the guard above works before and after the timetable phase.
    return set(inspect(session.get_bind()).get_table_names())


def _storage_headers() -> dict[str, str]:
    # Supabase Storage wants the service-role key both as a bearer token and as the apikey header.
    key = get_settings().supabase_service_role_key
    return {"Authorization": f"Bearer {key}", "apikey": key}


def _storage_base_url() -> str:
    # Root of the Storage REST API for this project.
    return f"{get_settings().supabase_url.rstrip('/')}/storage/v1"


def ensure_bucket() -> None:
    # Creates the course-schemes bucket the first time something is uploaded; a second call is harmless.
    with httpx.Client(timeout=STORAGE_TIMEOUT_SECONDS) as client:
        existing = client.get(f"{_storage_base_url()}/bucket/{SCHEME_BUCKET}", headers=_storage_headers())
        if existing.status_code == 200:
            return

        created = client.post(
            f"{_storage_base_url()}/bucket",
            headers=_storage_headers(),
            json={"id": SCHEME_BUCKET, "name": SCHEME_BUCKET, "public": True},
        )
        if created.status_code >= 400 and "already exists" not in created.text.lower():
            raise StorageError(f"Could not create the {SCHEME_BUCKET} bucket: {created.text}")


def build_storage_path(program_id: int, scheme_year: int, filename: str) -> str:
    # Keeps uploads tidy per program and year, with a random prefix so re-uploads never collide.
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "-", PurePath(filename).name).strip("-") or "scheme"
    return f"{program_id}/{scheme_year}/{uuid.uuid4().hex[:8]}-{safe_name}"


def store_scheme_file(
    program_id: int, scheme_year: int, filename: str, data: bytes, content_type: str
) -> str:
    # Uploads the original document and hands back the public URL saved on the scheme row.
    ensure_bucket()
    path = build_storage_path(program_id, scheme_year, filename)

    with httpx.Client(timeout=STORAGE_TIMEOUT_SECONDS) as client:
        response = client.post(
            f"{_storage_base_url()}/object/{SCHEME_BUCKET}/{path}",
            headers={**_storage_headers(), "Content-Type": content_type or "application/octet-stream"},
            content=data,
        )

    if response.status_code >= 400:
        raise StorageError(f"Upload failed ({response.status_code}): {response.text}")

    return f"{_storage_base_url()}/object/public/{SCHEME_BUCKET}/{path}"


def remove_scheme_file(file_url: str) -> None:
    # Deletes a stored document. Used when a save fails after the upload already happened.
    marker = f"/object/public/{SCHEME_BUCKET}/"
    if marker not in file_url:
        return

    path = file_url.split(marker, 1)[1]
    with httpx.Client(timeout=STORAGE_TIMEOUT_SECONDS) as client:
        client.delete(f"{_storage_base_url()}/object/{SCHEME_BUCKET}/{path}", headers=_storage_headers())
