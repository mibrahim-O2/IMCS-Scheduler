"""Render a saved Timetable to Word (.docx) with the IMCS logo in the header.

See docs/PROJECT_ARCHITECTURE.md §9. PDF export is not built in Phase 9 see
docs/PROJECT_AUDIT.md for what's still open.
"""

import io
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Classroom, Course, Division, Teacher, Timetable, TimetableSession
from app.scheduler.chromosome import DAYS

LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "imcs-logo.png"


def build_timetable_document(session: Session, timetable: Timetable) -> bytes:
    # Turns one saved Timetable into a .docx file in memory: a title page line, then one
    # heading + table per division, each table a plain weekly grid (Day, Time, Course,
    # Teacher, Room) in the same day order the frontend shows (docs/PROJECT_ARCHITECTURE.md
    # §9's "general/continuous" layout, not a fixed day x period grid).
    document = Document()
    _add_logo_header(document)

    title = document.add_heading(timetable.label, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_line = document.add_paragraph(
        f"Status: {timetable.status.value} · "
        f"{'Converged' if timetable.converged else 'Not converged'} · "
        f"Generated {timetable.created_at.strftime('%d %b %Y, %H:%M')}"
    )
    status_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for _division_id, division_label, rows in _collect_schedule(session, timetable.id):
        document.add_heading(division_label, level=2)
        if not rows:
            document.add_paragraph("No sessions scheduled.")
            continue

        table = document.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        for cell, text in zip(header_cells, ("Day", "Time", "Course", "Teacher", "Room"), strict=True):
            cell.text = text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for row in rows:
            cells = table.add_row().cells
            cells[0].text = row["day"]
            cells[1].text = f"{row['start_time']}–{row['end_time']}"
            course_text = f"{row['course_code']} {row['course_name']}"
            if row["is_lab"]:
                course_text += " (Lab)"
            cells[2].text = course_text
            cells[3].text = row["teacher_name"]
            cells[4].text = row["room_name"]

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_logo_header(document: Document) -> None:
    # Places the IMCS crest at the top of every page. Falls back to a plain text header if
    # the logo file is somehow missing, rather than failing the whole export over an image.
    header = document.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_PATH.exists():
        paragraph.add_run().add_picture(str(LOGO_PATH), width=Inches(0.6))
    else:
        run = paragraph.add_run("IMCS Scheduler")
        run.bold = True
        run.font.size = Pt(14)


def _collect_schedule(session: Session, timetable_id: int) -> list[tuple[int, str, list[dict]]]:
    # Reads the saved sessions straight from the database and groups them the same way the
    # /timetables/{id} API response does (division -> day, in Mon-Fri order, each day sorted
    # by start time) kept as a self-contained read here rather than imported from the
    # endpoint layer, since this is a plain query with no scheduling logic to share.
    sessions = session.scalars(
        select(TimetableSession).where(TimetableSession.timetable_id == timetable_id)
    ).all()
    divisions = {row.id: row for row in session.scalars(select(Division))}
    courses = {row.id: row for row in session.scalars(select(Course))}
    teachers = {row.id: row for row in session.scalars(select(Teacher))}
    rooms = {row.id: row for row in session.scalars(select(Classroom))}

    by_division: dict[int, list[TimetableSession]] = defaultdict(list)
    for row in sessions:
        for division_id in row.division_ids:
            by_division[division_id].append(row)

    result = []
    for division_id in sorted(by_division, key=lambda d: divisions[d].label if d in divisions else ""):
        division = divisions.get(division_id)
        label = division.label if division else f"Division {division_id}"
        ordered = sorted(by_division[division_id], key=lambda r: (DAYS.index(r.day), r.start_time))
        rows = [
            {
                "day": row.day,
                "start_time": row.start_time.strftime("%H:%M"),
                "end_time": row.end_time.strftime("%H:%M"),
                "course_code": courses[row.course_id].code,
                "course_name": courses[row.course_id].name,
                "teacher_name": teachers[row.teacher_id].full_name,
                "room_name": rooms[row.room_id].name,
                "is_lab": row.is_lab,
            }
            for row in ordered
        ]
        result.append((division_id, label, rows))
    return result
